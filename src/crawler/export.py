"""Optional Excel export from SQLite."""

from __future__ import annotations

from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from src.config import ROOT, get_config
from src.db import list_all_jobs


def export_excel(path: str | None = None, user_id: int | None = None) -> str:
    from src.settings import applications_dir, user_data_dir
    from src.tenant import resolve_user_id

    uid = resolve_user_id(user_id)
    out = path or str(user_data_dir(uid) / "jobs_results.xlsx")
    jobs = list_all_jobs(limit=10000, user_id=uid)

    wb = Workbook()
    ws = wb.active
    ws.title = "All Jobs"
    COLS = [
        "Source", "Title", "Company", "Location", "Salary", "Tags",
        "Date Posted", "URL", "Description", "Match Score", "Status", "Notes",
    ]
    accent = "4859FF"
    header_fill = PatternFill("solid", start_color=accent)
    header_font = Font(bold=True, color="FFFFFF", size=10, name="Arial")
    thin = Side(style="thin", color="D5D8DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_i, col_name in enumerate(COLS, 1):
        c = ws.cell(row=1, column=col_i, value=col_name)
        c.font = header_font
        c.fill = header_fill
        c.border = border

    for row_i, job in enumerate(jobs, 2):
        row_data = [
            job.get("source"), job.get("title"), job.get("company"),
            job.get("location"), job.get("salary"), job.get("tags"),
            job.get("date_posted"), job.get("url"),
            job.get("description_full") or job.get("description_short"),
            job.get("match_score"), job.get("status"), job.get("notes"),
        ]
        for col_i, val in enumerate(row_data, 1):
            c = ws.cell(row=row_i, column=col_i, value=str(val) if val else "")
            c.border = border

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Job Crawl Summary"
    ws2["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    counts = Counter(j.get("source", "") for j in jobs)
    for row_i, (src, cnt) in enumerate(sorted(counts.items()), 4):
        ws2.cell(row=row_i, column=1, value=src)
        ws2.cell(row=row_i, column=2, value=cnt)

    wb.save(out)
    print(f"Exported → {out}")
    return out
